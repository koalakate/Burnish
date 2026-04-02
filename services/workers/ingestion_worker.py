"""Ingestion worker — parses uploaded PPTX to CSM and generates thumbnails.

Job data:
    deck_id: str  — UUID of the Deck row
    org_id: str   — UUID of the owning organization
"""

from __future__ import annotations

import io
import logging
import tempfile
import uuid
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation

from services.db.engine import async_session
from services.db.models.deck import Deck, DeckStatus
from services.ingestion.pptx_parser import parse_pptx
from services.storage.r2 import R2Client

logger = logging.getLogger(__name__)

# Thumbnail size (width in pixels; height follows aspect ratio)
THUMBNAIL_WIDTH = 800


def _generate_thumbnails(pptx_path: Path, slide_count: int) -> list[bytes]:
    """Generate placeholder slide thumbnails as PNGs.

    python-pptx cannot render slides to images natively.  We create simple
    coloured placeholder images; a real pipeline would use LibreOffice or
    a headless renderer.  This keeps the worker functional end-to-end.
    """
    prs = Presentation(str(pptx_path))
    width_emu = prs.slide_width or 12192000
    height_emu = prs.slide_height or 6858000
    aspect = height_emu / width_emu
    thumb_w = THUMBNAIL_WIDTH
    thumb_h = int(thumb_w * aspect)

    thumbnails: list[bytes] = []
    for idx in range(slide_count):
        img = Image.new("RGB", (thumb_w, thumb_h), color=(245, 245, 245))
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        thumbnails.append(buf.getvalue())
    return thumbnails


async def process_ingestion_job(job: Any, _token: Any = None) -> dict[str, Any]:
    """Process an ingestion job: download PPTX, parse to CSM, upload thumbnails."""
    data = job.data if hasattr(job, "data") else job
    deck_id = data["deck_id"]
    org_id = data["org_id"]
    logger.info("Ingestion job started: deck_id=%s org_id=%s", deck_id, org_id)

    r2 = R2Client()

    async with async_session() as db:
        # Update status to parsing
        deck = await db.get(Deck, uuid.UUID(deck_id))
        if not deck:
            raise ValueError(f"Deck {deck_id} not found")
        deck.status = DeckStatus.parsing
        await db.commit()

        try:
            # Download PPTX from R2
            pptx_bytes = r2.download_file(deck.source_ref)

            # Parse to CSM
            with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
                tmp.write(pptx_bytes)
                tmp_path = Path(tmp.name)

            csm = parse_pptx(tmp_path)

            # Generate thumbnails
            thumbnails = _generate_thumbnails(tmp_path, len(csm.slides))

            # Clean up temp file
            tmp_path.unlink(missing_ok=True)

            # Store CSM as JSON in R2
            csm_key = r2.org_key(org_id, f"decks/{deck_id}/csm.json")
            csm_json = csm.model_dump_json()
            r2.upload_file(csm_json.encode(), csm_key, content_type="application/json")

            # Upload thumbnails to R2
            thumbnail_refs: list[str] = []
            for idx, thumb_bytes in enumerate(thumbnails):
                thumb_key = r2.org_key(org_id, f"decks/{deck_id}/thumbnails/slide_{idx}.png")
                r2.upload_file(thumb_bytes, thumb_key, content_type="image/png")
                thumbnail_refs.append(thumb_key)

            # Update deck record
            deck.csm_ref = csm_key
            deck.slide_count = len(csm.slides)
            deck.status = DeckStatus.parsed
            await db.commit()

            logger.info(
                "Ingestion complete: deck_id=%s slides=%d", deck_id, len(csm.slides)
            )
            return {
                "deck_id": deck_id,
                "slide_count": len(csm.slides),
                "csm_ref": csm_key,
                "thumbnail_refs": thumbnail_refs,
            }

        except Exception:
            logger.exception("Ingestion failed for deck_id=%s", deck_id)
            deck.status = DeckStatus.failed
            await db.commit()
            raise
