"""Integration tests for BullMQ workers.

Tests the full upload → check → correction flow by mocking R2 and DB,
but using real worker processing functions.
"""

from __future__ import annotations

import io
import uuid
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from pptx import Presentation

from packages.csm.brand import (
    BrandColor,
    BrandFont,
    BrandRuleset,
    LayoutRules,
    SizeRule,
)
from packages.csm.models import (
    CSM,
    BoundingBox,
    Color,
    Font,
    Paragraph,
    Slide,
    SlideBackground,
    TextElement,
    TextRun,
)
from services.db.models.brand import BrandRuleset as BrandRulesetModel
from services.db.models.check import (
    CheckRun,
    CheckStatus,
    CorrectionStatus,
    SlideCheckResult,
)
from services.db.models.check import (
    Issue as IssueModel,
)
from services.db.models.check import (
    Severity as DBSeverity,
)
from services.db.models.deck import Deck, DeckStatus
from services.workers.check_worker import process_check_job
from services.workers.correction_worker import process_correction_job
from services.workers.ingestion_worker import process_ingestion_job

# ---------------------------------------------------------------------------
# Test data helpers
# ---------------------------------------------------------------------------

ORG_ID = str(uuid.uuid4())
USER_ID = str(uuid.uuid4())
DECK_ID = str(uuid.uuid4())
CHECK_RUN_ID = str(uuid.uuid4())
RULESET_ID = str(uuid.uuid4())


def _create_test_pptx() -> bytes:
    """Create a minimal PPTX file for testing."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    if title and title.text_frame:
        title.text_frame.text = "Test Slide"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_brand_ruleset() -> BrandRuleset:
    """Create a basic brand ruleset for testing."""
    return BrandRuleset(
        name="Test Brand",
        colors=[
            BrandColor(name="Primary", hex="#0066CC", r=0, g=102, b=204),
            BrandColor(name="White", hex="#FFFFFF", r=255, g=255, b=255),
        ],
        fonts=[
            BrandFont(family="Arial"),
        ],
        size_rules=[
            SizeRule(role="title", min_pt=24, max_pt=44),
            SizeRule(role="body", min_pt=14, max_pt=24),
        ],
        layout_rules=LayoutRules(),
    )


def _make_csm() -> CSM:
    """Create a simple CSM for testing."""
    return CSM(
        title="Test Deck",
        author="test",
        width=12192000,
        height=6858000,
        slides=[
            Slide(
                index=0,
                elements=[
                    TextElement(
                        id="elem_1",
                        bbox=BoundingBox(x=100, y=100, width=500, height=100),
                        paragraphs=[
                            Paragraph(
                                runs=[
                                    TextRun(
                                        text="Hello World",
                                        font=Font(family="Arial", size_pt=18.0),
                                        color=Color(hex="#000000", r=0, g=0, b=0),
                                    )
                                ]
                            )
                        ],
                        role="title",
                    )
                ],
                background=SlideBackground(),
            )
        ],
    )


def _make_job(data: dict[str, Any]) -> MagicMock:
    """Create a mock BullMQ job object."""
    job = MagicMock()
    job.data = data
    return job


# ---------------------------------------------------------------------------
# Mock helpers
# ---------------------------------------------------------------------------


class MockR2:
    """In-memory R2 mock that stores files as a dict."""

    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}

    def upload_file(
        self, data: bytes, key: str, content_type: str = "application/octet-stream"
    ) -> str:
        self.files[key] = data
        return key

    def download_file(self, key: str) -> bytes:
        if key not in self.files:
            raise FileNotFoundError(f"R2 key not found: {key}")
        return self.files[key]

    def get_signed_url(self, key: str, expires_in: int = 900) -> str:
        return f"https://mock-r2.example.com/{key}?signed=true"

    def delete_file(self, key: str) -> None:
        self.files.pop(key, None)

    def org_key(self, org_id: str, key: str) -> str:
        return f"{org_id}/{key}"


# ---------------------------------------------------------------------------
# Ingestion worker tests
# ---------------------------------------------------------------------------


class TestIngestionWorker:
    @pytest.fixture()
    def r2(self) -> MockR2:
        mock_r2 = MockR2()
        pptx_bytes = _create_test_pptx()
        source_ref = f"{ORG_ID}/decks/{DECK_ID}/test.pptx"
        mock_r2.files[source_ref] = pptx_bytes
        return mock_r2

    @pytest.fixture()
    def deck(self) -> MagicMock:
        d = MagicMock(spec=Deck)
        d.id = uuid.UUID(DECK_ID)
        d.org_id = uuid.UUID(ORG_ID)
        d.source_ref = f"{ORG_ID}/decks/{DECK_ID}/test.pptx"
        d.status = DeckStatus.uploaded
        d.csm_ref = ""
        d.slide_count = 0
        return d

    @pytest.mark.asyncio()
    async def test_ingestion_parses_pptx_and_uploads_csm(
        self, r2: MockR2, deck: MagicMock
    ) -> None:
        mock_db = AsyncMock()
        mock_db.get = AsyncMock(return_value=deck)
        mock_db.commit = AsyncMock()

        mock_session_ctx = AsyncMock()
        mock_session_ctx.__aenter__ = AsyncMock(return_value=mock_db)
        mock_session_ctx.__aexit__ = AsyncMock(return_value=False)

        job = _make_job({"deck_id": DECK_ID, "org_id": ORG_ID})

        with (
            patch("services.workers.ingestion_worker.R2Client", return_value=r2),
            patch(
                "services.workers.ingestion_worker.async_session",
                return_value=mock_session_ctx,
            ),
        ):
            result = await process_ingestion_job(job)

        assert result["deck_id"] == DECK_ID
        assert result["slide_count"] >= 1
        assert result["csm_ref"] in r2.files
        assert deck.status == DeckStatus.parsed
        assert deck.slide_count >= 1

    @pytest.mark.asyncio()
    async def test_ingestion_sets_failed_on_error(self, deck: MagicMock) -> None:
        mock_db = AsyncMock()
        mock_db.get = AsyncMock(return_value=deck)
        mock_db.commit = AsyncMock()

        mock_session_ctx = AsyncMock()
        mock_session_ctx.__aenter__ = AsyncMock(return_value=mock_db)
        mock_session_ctx.__aexit__ = AsyncMock(return_value=False)

        # R2 that raises on download
        bad_r2 = MockR2()

        job = _make_job({"deck_id": DECK_ID, "org_id": ORG_ID})

        with (
            patch("services.workers.ingestion_worker.R2Client", return_value=bad_r2),
            patch(
                "services.workers.ingestion_worker.async_session",
                return_value=mock_session_ctx,
            ),
            pytest.raises(FileNotFoundError),
        ):
            await process_ingestion_job(job)

        assert deck.status == DeckStatus.failed


# ---------------------------------------------------------------------------
# Check worker tests
# ---------------------------------------------------------------------------


class TestCheckWorker:
    @pytest.fixture()
    def r2(self) -> MockR2:
        mock_r2 = MockR2()
        csm = _make_csm()
        csm_key = f"{ORG_ID}/decks/{DECK_ID}/csm.json"
        mock_r2.files[csm_key] = csm.model_dump_json().encode()
        return mock_r2

    @pytest.fixture()
    def deck(self) -> MagicMock:
        d = MagicMock(spec=Deck)
        d.id = uuid.UUID(DECK_ID)
        d.csm_ref = f"{ORG_ID}/decks/{DECK_ID}/csm.json"
        return d

    @pytest.fixture()
    def check_run(self) -> MagicMock:
        cr = MagicMock(spec=CheckRun)
        cr.id = uuid.UUID(CHECK_RUN_ID)
        cr.deck_id = uuid.UUID(DECK_ID)
        cr.ruleset_id = uuid.UUID(RULESET_ID)
        cr.org_id = uuid.UUID(ORG_ID)
        cr.status = CheckStatus.queued
        cr.started_at = None
        cr.completed_at = None
        cr.dqs_overall = None
        cr.issue_count_error = 0
        cr.issue_count_warning = 0
        cr.issue_count_info = 0
        cr.exported_pptx_ref = None
        return cr

    @pytest.fixture()
    def ruleset(self) -> MagicMock:
        brand = _make_brand_ruleset()
        rs = MagicMock(spec=BrandRulesetModel)
        rs.id = uuid.UUID(RULESET_ID)
        rs.rules = brand.model_dump()
        return rs

    @pytest.mark.asyncio()
    async def test_check_runs_evaluators_and_stores_results(
        self,
        r2: MockR2,
        deck: MagicMock,
        check_run: MagicMock,
        ruleset: MagicMock,
    ) -> None:
        mock_db = AsyncMock()

        def mock_get(model: type, id: uuid.UUID) -> Any:
            if model is CheckRun:
                return check_run
            if model is Deck:
                return deck
            if model is BrandRulesetModel:
                return ruleset
            return None

        mock_db.get = AsyncMock(side_effect=mock_get)
        mock_db.add = MagicMock()
        mock_db.flush = AsyncMock()
        mock_db.commit = AsyncMock()

        mock_session_ctx = AsyncMock()
        mock_session_ctx.__aenter__ = AsyncMock(return_value=mock_db)
        mock_session_ctx.__aexit__ = AsyncMock(return_value=False)

        job = _make_job({
            "check_run_id": CHECK_RUN_ID,
            "deck_id": DECK_ID,
            "org_id": ORG_ID,
        })

        with (
            patch("services.workers.check_worker.R2Client", return_value=r2),
            patch(
                "services.workers.check_worker.async_session",
                return_value=mock_session_ctx,
            ),
        ):
            result = await process_check_job(job)

        assert result["check_run_id"] == CHECK_RUN_ID
        assert isinstance(result["dqs_overall"], float)
        assert check_run.status == CheckStatus.complete
        assert check_run.dqs_overall is not None
        assert check_run.completed_at is not None
        assert mock_db.add.called  # Slide results and issues were added

    @pytest.mark.asyncio()
    async def test_check_marks_failed_on_error(
        self, check_run: MagicMock
    ) -> None:
        mock_db = AsyncMock()
        mock_db.get = AsyncMock(side_effect=[check_run, None])  # deck returns None
        mock_db.commit = AsyncMock()

        mock_session_ctx = AsyncMock()
        mock_session_ctx.__aenter__ = AsyncMock(return_value=mock_db)
        mock_session_ctx.__aexit__ = AsyncMock(return_value=False)

        job = _make_job({
            "check_run_id": CHECK_RUN_ID,
            "deck_id": DECK_ID,
            "org_id": ORG_ID,
        })

        with (
            patch("services.workers.check_worker.R2Client", return_value=MockR2()),
            patch(
                "services.workers.check_worker.async_session",
                return_value=mock_session_ctx,
            ),
            pytest.raises(ValueError, match="not found"),
        ):
            await process_check_job(job)

        assert check_run.status == CheckStatus.failed


# ---------------------------------------------------------------------------
# Correction worker tests
# ---------------------------------------------------------------------------


class TestCorrectionWorker:
    @pytest.fixture()
    def r2(self) -> MockR2:
        mock_r2 = MockR2()
        # Store CSM
        csm = _make_csm()
        csm_key = f"{ORG_ID}/decks/{DECK_ID}/csm.json"
        mock_r2.files[csm_key] = csm.model_dump_json().encode()
        # Store original PPTX
        pptx_bytes = _create_test_pptx()
        source_ref = f"{ORG_ID}/decks/{DECK_ID}/test.pptx"
        mock_r2.files[source_ref] = pptx_bytes
        return mock_r2

    @pytest.fixture()
    def deck(self) -> MagicMock:
        d = MagicMock(spec=Deck)
        d.id = uuid.UUID(DECK_ID)
        d.csm_ref = f"{ORG_ID}/decks/{DECK_ID}/csm.json"
        d.source_ref = f"{ORG_ID}/decks/{DECK_ID}/test.pptx"
        return d

    @pytest.fixture()
    def check_run_with_issues(self) -> MagicMock:
        cr = MagicMock(spec=CheckRun)
        cr.id = uuid.UUID(CHECK_RUN_ID)
        cr.deck_id = uuid.UUID(DECK_ID)
        cr.ruleset_id = uuid.UUID(RULESET_ID)
        cr.org_id = uuid.UUID(ORG_ID)
        cr.status = CheckStatus.complete
        cr.exported_pptx_ref = None

        # Create mock slide result with issues
        sr = MagicMock(spec=SlideCheckResult)
        sr.slide_index = 0
        issue = MagicMock(spec=IssueModel)
        issue.id = uuid.uuid4()
        issue.rule_type = "typography"
        issue.severity = DBSeverity.warning
        issue.message = "Font not in brand palette"
        issue.element_id = "elem_1"
        issue.original_value = "Comic Sans"
        issue.expected_value = "Arial"
        issue.correction_status = CorrectionStatus.accepted
        sr.issues = [issue]
        cr.slide_results = [sr]
        return cr

    @pytest.fixture()
    def ruleset(self) -> MagicMock:
        brand = _make_brand_ruleset()
        rs = MagicMock(spec=BrandRulesetModel)
        rs.id = uuid.UUID(RULESET_ID)
        rs.rules = brand.model_dump()
        return rs

    @pytest.mark.asyncio()
    async def test_correction_exports_pptx(
        self,
        r2: MockR2,
        deck: MagicMock,
        check_run_with_issues: MagicMock,
        ruleset: MagicMock,
    ) -> None:
        mock_db = AsyncMock()

        def mock_get(model: type, id: uuid.UUID) -> Any:
            if model is Deck:
                return deck
            if model is BrandRulesetModel:
                return ruleset
            return None

        mock_db.get = AsyncMock(side_effect=mock_get)
        mock_db.commit = AsyncMock()

        # Mock the selectinload query
        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = check_run_with_issues
        mock_db.execute = AsyncMock(return_value=mock_result)

        mock_session_ctx = AsyncMock()
        mock_session_ctx.__aenter__ = AsyncMock(return_value=mock_db)
        mock_session_ctx.__aexit__ = AsyncMock(return_value=False)

        job = _make_job({
            "check_run_id": CHECK_RUN_ID,
            "deck_id": DECK_ID,
            "org_id": ORG_ID,
        })

        with (
            patch("services.workers.correction_worker.R2Client", return_value=r2),
            patch(
                "services.workers.correction_worker.async_session",
                return_value=mock_session_ctx,
            ),
        ):
            result = await process_correction_job(job)

        assert result["check_run_id"] == CHECK_RUN_ID
        assert result["exported_pptx_ref"] is not None
        assert check_run_with_issues.exported_pptx_ref is not None

        # Verify the corrected PPTX was uploaded to R2
        export_key = result["exported_pptx_ref"]
        assert export_key in r2.files
        assert len(r2.files[export_key]) > 0  # Non-empty PPTX bytes


# ---------------------------------------------------------------------------
# End-to-end flow test
# ---------------------------------------------------------------------------


class TestEndToEndFlow:
    """Test the full upload → check → fix-all → export flow."""

    @pytest.mark.asyncio()
    async def test_full_flow(self) -> None:
        r2 = MockR2()

        # Step 1: Setup - upload PPTX to R2 (simulating the upload API)
        pptx_bytes = _create_test_pptx()
        source_ref = f"{ORG_ID}/decks/{DECK_ID}/test.pptx"
        r2.files[source_ref] = pptx_bytes

        deck = MagicMock(spec=Deck)
        deck.id = uuid.UUID(DECK_ID)
        deck.org_id = uuid.UUID(ORG_ID)
        deck.source_ref = source_ref
        deck.status = DeckStatus.uploaded
        deck.csm_ref = ""
        deck.slide_count = 0

        # Step 2: Run ingestion
        mock_db_ingest = AsyncMock()
        mock_db_ingest.get = AsyncMock(return_value=deck)
        mock_db_ingest.commit = AsyncMock()

        mock_session_ingest = AsyncMock()
        mock_session_ingest.__aenter__ = AsyncMock(return_value=mock_db_ingest)
        mock_session_ingest.__aexit__ = AsyncMock(return_value=False)

        ingest_job = _make_job({"deck_id": DECK_ID, "org_id": ORG_ID})

        with (
            patch("services.workers.ingestion_worker.R2Client", return_value=r2),
            patch(
                "services.workers.ingestion_worker.async_session",
                return_value=mock_session_ingest,
            ),
        ):
            ingest_result = await process_ingestion_job(ingest_job)

        assert ingest_result["slide_count"] >= 1
        assert deck.status == DeckStatus.parsed

        # Step 3: Run check
        brand = _make_brand_ruleset()
        ruleset = MagicMock(spec=BrandRulesetModel)
        ruleset.id = uuid.UUID(RULESET_ID)
        ruleset.rules = brand.model_dump()

        check_run = MagicMock(spec=CheckRun)
        check_run.id = uuid.UUID(CHECK_RUN_ID)
        check_run.deck_id = uuid.UUID(DECK_ID)
        check_run.ruleset_id = uuid.UUID(RULESET_ID)
        check_run.org_id = uuid.UUID(ORG_ID)
        check_run.status = CheckStatus.queued
        check_run.started_at = None
        check_run.completed_at = None
        check_run.dqs_overall = None
        check_run.issue_count_error = 0
        check_run.issue_count_warning = 0
        check_run.issue_count_info = 0

        mock_db_check = AsyncMock()

        def check_get(model: type, id: uuid.UUID) -> Any:
            if model is CheckRun:
                return check_run
            if model is Deck:
                return deck
            if model is BrandRulesetModel:
                return ruleset
            return None

        mock_db_check.get = AsyncMock(side_effect=check_get)
        mock_db_check.add = MagicMock()
        mock_db_check.flush = AsyncMock()
        mock_db_check.commit = AsyncMock()

        mock_session_check = AsyncMock()
        mock_session_check.__aenter__ = AsyncMock(return_value=mock_db_check)
        mock_session_check.__aexit__ = AsyncMock(return_value=False)

        check_job = _make_job({
            "check_run_id": CHECK_RUN_ID,
            "deck_id": DECK_ID,
            "org_id": ORG_ID,
        })

        with (
            patch("services.workers.check_worker.R2Client", return_value=r2),
            patch(
                "services.workers.check_worker.async_session",
                return_value=mock_session_check,
            ),
        ):
            check_result = await process_check_job(check_job)

        assert check_run.status == CheckStatus.complete
        assert check_result["dqs_overall"] is not None

        # Step 4: Run correction (simulate fix-all)
        # Create mock check run with the issues that were stored
        cr_for_correction = MagicMock(spec=CheckRun)
        cr_for_correction.id = uuid.UUID(CHECK_RUN_ID)
        cr_for_correction.deck_id = uuid.UUID(DECK_ID)
        cr_for_correction.ruleset_id = uuid.UUID(RULESET_ID)
        cr_for_correction.org_id = uuid.UUID(ORG_ID)
        cr_for_correction.status = CheckStatus.complete
        cr_for_correction.exported_pptx_ref = None

        sr = MagicMock(spec=SlideCheckResult)
        sr.slide_index = 0
        sr.issues = []  # No issues to correct in this minimal test
        cr_for_correction.slide_results = [sr]

        mock_db_correct = AsyncMock()

        def correct_get(model: type, id: uuid.UUID) -> Any:
            if model is Deck:
                return deck
            if model is BrandRulesetModel:
                return ruleset
            return None

        mock_db_correct.get = AsyncMock(side_effect=correct_get)
        mock_db_correct.commit = AsyncMock()

        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = cr_for_correction
        mock_db_correct.execute = AsyncMock(return_value=mock_result)

        mock_session_correct = AsyncMock()
        mock_session_correct.__aenter__ = AsyncMock(return_value=mock_db_correct)
        mock_session_correct.__aexit__ = AsyncMock(return_value=False)

        correction_job = _make_job({
            "check_run_id": CHECK_RUN_ID,
            "deck_id": DECK_ID,
            "org_id": ORG_ID,
        })

        with (
            patch("services.workers.correction_worker.R2Client", return_value=r2),
            patch(
                "services.workers.correction_worker.async_session",
                return_value=mock_session_correct,
            ),
        ):
            correction_result = await process_correction_job(correction_job)

        # Verify the corrected PPTX was exported
        assert correction_result["exported_pptx_ref"] is not None
        export_key = correction_result["exported_pptx_ref"]
        assert export_key in r2.files

        # Verify the exported file is a valid PPTX (can be opened)
        from pptx import Presentation as PrsCheck

        pptx_data = io.BytesIO(r2.files[export_key])
        prs = PrsCheck(pptx_data)
        assert len(prs.slides) >= 1


# ---------------------------------------------------------------------------
# Queue helper tests
# ---------------------------------------------------------------------------


class TestQueueHelper:
    @pytest.mark.asyncio()
    async def test_enqueue_job_creates_queue_and_adds_job(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        # Restore real enqueue_job (overriding the autouse mock)
        import importlib

        import services.workers.queue as queue_mod

        importlib.reload(queue_mod)

        mock_queue = AsyncMock()
        mock_job = MagicMock()
        mock_job.id = "test-job-123"
        mock_queue.add = AsyncMock(return_value=mock_job)
        mock_queue.close = AsyncMock()

        with patch(
            "services.workers.queue.Queue", return_value=mock_queue
        ):
            job_id = await queue_mod.enqueue_job("check", {"check_run_id": "abc"})

        assert job_id == "test-job-123"
        mock_queue.add.assert_called_once_with("check", {"check_run_id": "abc"})
