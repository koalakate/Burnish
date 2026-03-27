"""Tests for PPTX parser — uses a hand-crafted PPTX fixture."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from packages.csm.models import (
    ShapeElement,
    TableElement,
    TextElement,
)
from services.ingestion.pptx_parser import parse_pptx
from services.ingestion.thumbnail import generate_deck_thumbnails, generate_thumbnail


def _create_test_pptx(path: Path) -> None:
    """Create a minimal PPTX with various shape types for testing."""
    prs = Presentation()
    prs.slide_width = Emu(9144000)   # 10 inches
    prs.slide_height = Emu(6858000)  # 7.5 inches
    prs.core_properties.title = "Test Presentation"
    prs.core_properties.author = "Test Author"

    # Slide 1: Title slide with text
    slide_layout = prs.slide_layouts[0]  # Title Slide
    slide1 = prs.slides.add_slide(slide_layout)
    title = slide1.shapes.title
    title.text = "Hello World"
    subtitle = slide1.placeholders[1]
    subtitle.text = "A subtitle"

    # Slide 2: Content with text box, auto-shape, and table
    blank_layout = prs.slide_layouts[6]  # Blank
    slide2 = prs.slides.add_slide(blank_layout)

    # Add text box with formatted text
    tx_box = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Formatted text"
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    # Add a second run in the same paragraph
    run2 = p.add_run()
    run2.text = " with italic"
    run2.font.name = "Arial"
    run2.font.size = Pt(14)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)

    # Add auto-shape (rectangle)
    shape = slide2.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(3), Inches(1.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x00, 0xCC, 0x00)

    # Add table
    table_shape = slide2.shapes.add_table(2, 3, Inches(5), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(0, 2).text = "Header 3"
    table.cell(1, 0).text = "Row 1"
    table.cell(1, 1).text = "Data"
    table.cell(1, 2).text = "More data"

    # Slide 3: Slide with notes
    slide3 = prs.slides.add_slide(blank_layout)
    tx_box3 = slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    tf3 = tx_box3.text_frame
    tf3.text = "Slide with notes"
    notes_slide = slide3.notes_slide
    notes_slide.notes_text_frame.text = "These are speaker notes."

    prs.save(str(path))


@pytest.fixture
def test_pptx(tmp_path: Path) -> Path:
    """Create a temporary test PPTX file."""
    pptx_path = tmp_path / "test.pptx"
    _create_test_pptx(pptx_path)
    return pptx_path


class TestParsePptx:
    """Tests for the parse_pptx function."""

    def test_slide_count(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        assert len(csm.slides) == 3

    def test_presentation_dimensions(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        assert csm.width == 9144000.0
        assert csm.height == 6858000.0

    def test_presentation_metadata(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        assert csm.title == "Test Presentation"
        assert csm.author == "Test Author"

    def test_slide_1_has_text_elements(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[0]
        text_elements = [e for e in slide.elements if isinstance(e, TextElement)]
        assert len(text_elements) >= 2  # title + subtitle

    def test_slide_1_title_text(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[0]
        text_elements = [e for e in slide.elements if isinstance(e, TextElement)]
        texts = []
        for te in text_elements:
            for p in te.paragraphs:
                for r in p.runs:
                    texts.append(r.text)
        assert "Hello World" in texts

    def test_slide_2_element_types(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        types = {type(e).__name__ for e in slide.elements}
        assert "TextElement" in types
        assert "ShapeElement" in types
        assert "TableElement" in types

    def test_text_font_properties(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        text_elements = [e for e in slide.elements if isinstance(e, TextElement)]
        for te in text_elements:
            for p in te.paragraphs:
                for r in p.runs:
                    if r.text == "Formatted text":
                        assert r.font.family == "Arial"
                        assert r.font.size_pt is not None
                        assert abs(r.font.size_pt - 18.0) < 0.1
                        assert r.font.weight == 700  # bold
                        return
        pytest.fail("Could not find 'Formatted text' run")

    def test_text_color_extraction(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        text_elements = [e for e in slide.elements if isinstance(e, TextElement)]
        for te in text_elements:
            for p in te.paragraphs:
                for r in p.runs:
                    if r.text == "Formatted text":
                        assert r.color is not None
                        assert r.color.r == 255
                        assert r.color.g == 0
                        assert r.color.b == 0
                        return
        pytest.fail("Could not find 'Formatted text' run with color")

    def test_italic_run(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        text_elements = [e for e in slide.elements if isinstance(e, TextElement)]
        for te in text_elements:
            for p in te.paragraphs:
                for r in p.runs:
                    if "italic" in r.text:
                        assert r.font.italic is True
                        assert r.color is not None
                        assert r.color.b == 255
                        return
        pytest.fail("Could not find italic run")

    def test_shape_fill_color(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        shapes = [e for e in slide.elements if isinstance(e, ShapeElement)]
        assert len(shapes) >= 1
        green_shape = shapes[0]
        assert green_shape.fill_color is not None
        assert green_shape.fill_color.g == 204  # 0xCC

    def test_table_structure(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        tables = [e for e in slide.elements if isinstance(e, TableElement)]
        assert len(tables) == 1
        tbl = tables[0]
        assert tbl.rows == 2
        assert tbl.cols == 3
        assert len(tbl.cells) == 6  # 2x3

    def test_table_cell_text(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[1]
        tables = [e for e in slide.elements if isinstance(e, TableElement)]
        tbl = tables[0]
        cell_texts = []
        for cell in tbl.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    cell_texts.append(r.text)
        assert "Header 1" in cell_texts
        assert "Row 1" in cell_texts

    def test_speaker_notes(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[2]
        assert slide.notes is not None
        assert "speaker notes" in slide.notes.lower()

    def test_bounding_boxes_nonzero(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        for slide in csm.slides:
            for element in slide.elements:
                bbox = element.bbox
                if bbox.width > 0:
                    assert bbox.width > 0
                    assert bbox.height > 0

    def test_layout_name(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        assert csm.slides[0].layout_name is not None

    def test_csm_json_roundtrip(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        json_str = csm.model_dump_json()
        csm2 = type(csm).model_validate_json(json_str)
        assert len(csm2.slides) == len(csm.slides)
        assert csm2.width == csm.width
        assert csm2.title == csm.title


class TestThumbnailGeneration:
    """Tests for the thumbnail generator."""

    def test_single_thumbnail(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        slide = csm.slides[0]
        png_bytes = generate_thumbnail(slide, csm.width, csm.height)
        assert len(png_bytes) > 0
        assert png_bytes[:4] == b"\x89PNG"

    def test_deck_thumbnails(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        thumbnails = generate_deck_thumbnails(csm)
        assert len(thumbnails) == 3
        for thumb in thumbnails:
            assert thumb[:4] == b"\x89PNG"

    def test_thumbnail_dimensions(self, test_pptx: Path) -> None:
        csm = parse_pptx(test_pptx)
        png_bytes = generate_thumbnail(csm.slides[0], csm.width, csm.height, 800, 600)
        img = Image.open(BytesIO(png_bytes))
        assert img.size == (800, 600)
