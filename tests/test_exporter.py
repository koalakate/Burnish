"""Tests for the PPTX exporter — corrected CSM -> .pptx -> re-parse -> verify."""

from __future__ import annotations

import io
import tempfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from packages.csm.models import Color
from services.correction.exporter import export_pptx
from services.ingestion.pptx_parser import parse_pptx

# ---------------------------------------------------------------------------
# Helpers — build minimal PPTX files for round-trip tests
# ---------------------------------------------------------------------------


def _make_pptx_with_text(
    text: str = "Hello",
    font_name: str = "Calibri",
    font_size_pt: float = 18.0,
    font_color_rgb: tuple[int, int, int] = (0, 0, 0),
    bold: bool = False,
) -> Path:
    """Create a minimal PPTX with one text box and return the temp file path."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[6]  # blank layout
    slide = prs.slides.add_slide(slide_layout)
    txbox = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1),
    )
    tf = txbox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.color.rgb = RGBColor(*font_color_rgb)
    run.font.bold = bold

    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    prs.save(tmp.name)
    tmp.close()
    return Path(tmp.name)


def _make_pptx_with_shape(
    fill_rgb: tuple[int, int, int] = (255, 0, 0),
) -> Path:
    """Create a minimal PPTX with one auto-shape and return the temp file path."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(1), Inches(1), Inches(3), Inches(2),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_rgb)

    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    prs.save(tmp.name)
    tmp.close()
    return Path(tmp.name)


def _make_pptx_with_table() -> Path:
    """Create a minimal PPTX with one 2x2 table."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    table_shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(4), Inches(2),
    )
    table = table_shape.table

    # Fill first cell with some text and color
    cell = table.cell(0, 0)
    cell.text = "A1"
    run = cell.text_frame.paragraphs[0].runs[0]
    run.font.name = "Calibri"
    run.font.size = Pt(12.0)
    run.font.color.rgb = RGBColor(255, 0, 0)

    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    prs.save(tmp.name)
    tmp.close()
    return Path(tmp.name)


def _make_pptx_with_notes() -> Path:
    """Create a PPTX with speaker notes on slide 1."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    txbox = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1),
    )
    tf = txbox.text_frame
    run = tf.paragraphs[0].add_run()
    run.text = "Content"
    run.font.name = "Calibri"
    run.font.size = Pt(18.0)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # Add speaker notes
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "These are my speaker notes"

    tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    prs.save(tmp.name)
    tmp.close()
    return Path(tmp.name)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


class TestExportTextCorrections:
    def test_font_family_change(self) -> None:
        """Changing font family in CSM is reflected in exported PPTX."""
        pptx_path = _make_pptx_with_text(font_name="Calibri")
        original_csm = parse_pptx(pptx_path)

        # Modify CSM: change font family
        corrected = original_csm.model_copy(deep=True)
        for slide in corrected.slides:
            for elem in slide.elements:
                if elem.type == "text":
                    for para in elem.paragraphs:
                        for run in para.runs:
                            run.font = run.font.model_copy(
                                update={"family": "Inter"}
                            )

        result_bytes = export_pptx(pptx_path, corrected)

        # Re-parse and verify
        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        text_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "text"
        ]
        assert len(text_elems) > 0
        for elem in text_elems:
            for para in elem.paragraphs:
                for run in para.runs:
                    assert run.font.family == "Inter"

    def test_font_size_change(self) -> None:
        """Changing font size in CSM is reflected in exported PPTX."""
        pptx_path = _make_pptx_with_text(font_size_pt=18.0)
        original_csm = parse_pptx(pptx_path)

        corrected = original_csm.model_copy(deep=True)
        for slide in corrected.slides:
            for elem in slide.elements:
                if elem.type == "text":
                    for para in elem.paragraphs:
                        for run in para.runs:
                            run.font = run.font.model_copy(
                                update={"size_pt": 24.0}
                            )

        result_bytes = export_pptx(pptx_path, corrected)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        text_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "text"
        ]
        assert len(text_elems) > 0
        for elem in text_elems:
            for para in elem.paragraphs:
                for run in para.runs:
                    assert run.font.size_pt == pytest.approx(24.0, abs=0.1)

    def test_text_color_change(self) -> None:
        """Changing text color in CSM is reflected in exported PPTX."""
        pptx_path = _make_pptx_with_text(font_color_rgb=(255, 0, 0))
        original_csm = parse_pptx(pptx_path)

        corrected = original_csm.model_copy(deep=True)
        new_color = Color(hex="#1B3A6B", r=27, g=58, b=107)
        for slide in corrected.slides:
            for elem in slide.elements:
                if elem.type == "text":
                    for para in elem.paragraphs:
                        for run in para.runs:
                            run.color = new_color

        result_bytes = export_pptx(pptx_path, corrected)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        text_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "text"
        ]
        assert len(text_elems) > 0
        for elem in text_elems:
            for para in elem.paragraphs:
                for run in para.runs:
                    assert run.color is not None
                    assert run.color.hex == "#1B3A6B"


class TestExportShapeCorrections:
    def test_shape_fill_color_change(self) -> None:
        """Changing shape fill color in CSM is reflected in exported PPTX."""
        pptx_path = _make_pptx_with_shape(fill_rgb=(255, 0, 0))
        original_csm = parse_pptx(pptx_path)

        corrected = original_csm.model_copy(deep=True)
        new_color = Color(hex="#1B3A6B", r=27, g=58, b=107)
        for slide in corrected.slides:
            for elem in slide.elements:
                if elem.type == "shape":
                    elem.fill_color = new_color

        result_bytes = export_pptx(pptx_path, corrected)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        shape_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "shape"
        ]
        assert len(shape_elems) > 0
        for elem in shape_elems:
            assert elem.fill_color is not None
            assert elem.fill_color.hex == "#1B3A6B"


class TestExportPositionCorrections:
    def test_element_repositioned(self) -> None:
        """Changing element position in CSM is reflected in exported PPTX."""
        pptx_path = _make_pptx_with_text()
        original_csm = parse_pptx(pptx_path)

        corrected = original_csm.model_copy(deep=True)
        new_x = int(Inches(2))
        new_y = int(Inches(2))
        for slide in corrected.slides:
            for elem in slide.elements:
                elem.bbox = elem.bbox.model_copy(
                    update={"x": float(new_x), "y": float(new_y)}
                )

        result_bytes = export_pptx(pptx_path, corrected)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        for slide in reparsed.slides:
            for elem in slide.elements:
                assert elem.bbox.x == pytest.approx(float(new_x), abs=1)
                assert elem.bbox.y == pytest.approx(float(new_y), abs=1)


class TestExportTableCorrections:
    def test_table_cell_text_color_change(self) -> None:
        """Changing table cell text color in CSM is reflected in exported PPTX."""
        pptx_path = _make_pptx_with_table()
        original_csm = parse_pptx(pptx_path)

        corrected = original_csm.model_copy(deep=True)
        new_color = Color(hex="#1B3A6B", r=27, g=58, b=107)
        for slide in corrected.slides:
            for elem in slide.elements:
                if elem.type == "table":
                    for cell in elem.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                run.color = new_color

        result_bytes = export_pptx(pptx_path, corrected)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        table_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "table"
        ]
        assert len(table_elems) > 0
        # Check cell (0,0) which had text
        cell_00 = next(
            c for c in table_elems[0].cells if c.row == 0 and c.col == 0
        )
        colored_runs = [
            r for p in cell_00.paragraphs for r in p.runs if r.color is not None
        ]
        assert len(colored_runs) > 0
        assert colored_runs[0].color is not None
        assert colored_runs[0].color.hex == "#1B3A6B"


class TestExportPreservesUntouched:
    def test_notes_preserved(self) -> None:
        """Speaker notes should survive the export round-trip."""
        pptx_path = _make_pptx_with_notes()
        original_csm = parse_pptx(pptx_path)

        # Don't modify the CSM at all
        result_bytes = export_pptx(pptx_path, original_csm)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        assert reparsed.slides[0].notes == "These are my speaker notes"

    def test_unmodified_elements_preserved(self) -> None:
        """Elements not in CSM corrections should remain unchanged."""
        pptx_path = _make_pptx_with_text(
            text="Hello", font_name="Calibri", font_size_pt=18.0,
        )
        original_csm = parse_pptx(pptx_path)

        # Export without any changes
        result_bytes = export_pptx(pptx_path, original_csm)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        text_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "text"
        ]
        assert len(text_elems) > 0
        run = text_elems[0].paragraphs[0].runs[0]
        assert run.text == "Hello"
        assert run.font.family == "Calibri"
        assert run.font.size_pt == pytest.approx(18.0, abs=0.1)

    def test_returns_valid_bytes(self) -> None:
        """The exported bytes should be a valid PPTX that python-pptx can open."""
        pptx_path = _make_pptx_with_text()
        csm = parse_pptx(pptx_path)
        result_bytes = export_pptx(pptx_path, csm)

        # Should be openable
        prs = Presentation(io.BytesIO(result_bytes))
        assert len(prs.slides) == 1


class TestExportMultipleCorrections:
    def test_combined_font_color_size(self) -> None:
        """Multiple corrections (font + color + size) applied together."""
        pptx_path = _make_pptx_with_text(
            text="Bad Slide",
            font_name="Comic Sans MS",
            font_size_pt=10.0,
            font_color_rgb=(255, 0, 0),
        )
        original_csm = parse_pptx(pptx_path)

        corrected = original_csm.model_copy(deep=True)
        new_color = Color(hex="#1B3A6B", r=27, g=58, b=107)
        for slide in corrected.slides:
            for elem in slide.elements:
                if elem.type == "text":
                    for para in elem.paragraphs:
                        for run in para.runs:
                            run.font = run.font.model_copy(
                                update={"family": "Inter", "size_pt": 24.0}
                            )
                            run.color = new_color

        result_bytes = export_pptx(pptx_path, corrected)

        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        tmp.write(result_bytes)
        tmp.close()
        reparsed = parse_pptx(Path(tmp.name))

        text_elems = [
            e for s in reparsed.slides for e in s.elements if e.type == "text"
        ]
        assert len(text_elems) > 0
        run = text_elems[0].paragraphs[0].runs[0]
        assert run.font.family == "Inter"
        assert run.font.size_pt == pytest.approx(24.0, abs=0.1)
        assert run.color is not None
        assert run.color.hex == "#1B3A6B"
